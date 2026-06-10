"""
SB30 부하 해리력 분석 - 1페이지 요약 보고서 PPT 생성
"""
import pandas as pd
import numpy as np
import os
import io
import warnings
warnings.filterwarnings('ignore')

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
plt.rcParams['font.family'] = 'Malgun Gothic'
plt.rcParams['axes.unicode_minus'] = False

from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BASE_DIR  = os.path.dirname(os.path.abspath(__file__))
XLS_FILE  = os.path.join(BASE_DIR, "분석결과_리포트.xlsx")
PPT_FILE  = os.path.join(BASE_DIR, "SB30_부하해리력_분석보고서.pptx")

# ── 색상 팔레트
CN  = RGBColor(0x1F, 0x39, 0x64)   # navy
CB  = RGBColor(0x26, 0x72, 0xC4)   # blue
CR  = RGBColor(0xC0, 0x00, 0x00)   # red
CO  = RGBColor(0xFF, 0x80, 0x00)   # orange
CG  = RGBColor(0x00, 0x70, 0x50)   # green
CW  = RGBColor(0xFF, 0xFF, 0xFF)   # white
CLG = RGBColor(0xF2, 0xF2, 0xF2)   # light gray
CMG = RGBColor(0xD0, 0xD0, 0xD0)   # mid gray
CDG = RGBColor(0x40, 0x40, 0x40)   # dark gray
CLB = RGBColor(0xBD, 0xD7, 0xEE)   # light blue
CMB = RGBColor(0x37, 0x5A, 0x8C)   # mid blue

# ══════════════════════════════════════════════════════
# 1. 데이터 로드
# ══════════════════════════════════════════════════════
df_trend = pd.read_excel(XLS_FILE, sheet_name='연도별추세')
df_car   = pd.read_excel(XLS_FILE, sheet_name='차종별NG')
df_pn    = pd.read_excel(XLS_FILE, sheet_name='품번별NG')
df_stage = pd.read_excel(XLS_FILE, sheet_name='단계별NG')
df_fc_a  = pd.read_excel(XLS_FILE, sheet_name='연도별예측')

df_spc_raw = pd.read_excel(XLS_FILE, sheet_name='SPC관리도', header=None, nrows=6)
spc_ucl    = float(df_spc_raw.iloc[1, 1])
spc_warn_u = float(df_spc_raw.iloc[2, 1])
spc_cl     = float(df_spc_raw.iloc[3, 1])
spc_warn_l = float(df_spc_raw.iloc[4, 1])
spc_lcl    = float(df_spc_raw.iloc[5, 1])

pred_2027 = df_fc_a[df_fc_a['연도'] == 2027][['하한(90%)', '상한(90%)']].values[0]
pred_2028 = df_fc_a[df_fc_a['연도'] == 2028][['하한(90%)', '상한(90%)']].values[0]

total_n  = int(df_trend['총건수'].sum())
total_ng = int(df_trend['NG전체'].sum())
ng_rate  = total_ng / total_n * 100

# ══════════════════════════════════════════════════════
# 2. matplotlib 연도별 추세 차트
# ══════════════════════════════════════════════════════
def make_trend_chart(w_in=5.6, h_in=3.8):
    fig, ax1 = plt.subplots(figsize=(w_in, h_in))
    fig.patch.set_facecolor('#F8F9FA')
    ax1.set_facecolor('#F8F9FA')

    years    = df_trend['_시트'].astype(str)
    ng_rates = df_trend['NG율(%)']
    s1_means = df_trend['S1_평균']
    x = range(len(years))

    colors = ['#C00000' if v > 15 else '#FF8000' if v > 5 else '#70AD47'
              for v in ng_rates]
    ax1.bar(x, ng_rates, color=colors, alpha=0.88, zorder=2, width=0.58)
    ax1.set_ylabel('NG율 (%)', color='#C00000', fontsize=9)
    ax1.tick_params(axis='y', labelcolor='#C00000', labelsize=8)
    ax1.set_xticks(x)
    ax1.set_xticklabels(years, rotation=38, ha='right', fontsize=8.5)
    ax1.set_ylim(0, max(ng_rates) * 1.5)
    ax1.grid(axis='y', alpha=0.35, zorder=0, linewidth=0.7)
    ax1.spines[['top', 'right']].set_visible(False)

    ax2 = ax1.twinx()
    ax2.plot(x, s1_means, 'o-', color='#2672C4', linewidth=2.2,
             markersize=5.5, zorder=3)
    for xi, yi in zip(x, s1_means):
        ax2.annotate(f'{yi:.1f}', (xi, yi), textcoords='offset points',
                     xytext=(0, 6), ha='center', fontsize=7, color='#2672C4')
    ax2.set_ylabel('S1 평균 (N)', color='#2672C4', fontsize=9)
    ax2.tick_params(axis='y', labelcolor='#2672C4', labelsize=8)
    ax2.set_ylim(35, 62)
    ax2.spines[['top']].set_visible(False)

    ax1.set_title('연도별 NG율(%)  vs  S1 평균(N)', fontsize=10.5,
                   fontweight='bold', color='#1F3964', pad=7)
    fig.tight_layout(pad=0.5)

    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=160, bbox_inches='tight',
                facecolor=fig.get_facecolor())
    buf.seek(0)
    plt.close(fig)
    return buf

# ══════════════════════════════════════════════════════
# 3. python-pptx 헬퍼
# ══════════════════════════════════════════════════════
def rect(slide, l, t, w, h, fill, line_clr=None, lw=0.5):
    shp = slide.shapes.add_shape(1, Cm(l), Cm(t), Cm(w), Cm(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line_clr:
        shp.line.color.rgb = line_clr
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    return shp

def txt(slide, l, t, w, h, text, sz=9, bold=False,
        clr=CDG, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Cm(l), Cm(t), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(sz)
    r.font.bold = bold
    r.font.color.rgb = clr
    return tb

def section_hdr(slide, l, t, w, title, color=CN):
    rect(slide, l, t, w, 0.58, color)
    txt(slide, l+0.15, t+0.1, w-0.3, 0.42, title,
        sz=9, bold=True, clr=CW)

def tbl_hdr(slide, l, t, w, cols):
    """cols: list of (label, col_x, col_w, align)"""
    rect(slide, l, t, w, 0.48, CMG)
    for label, cx, cw, al in cols:
        txt(slide, l+cx, t+0.07, cw, 0.36, label,
            sz=8, bold=True, clr=CN, align=al)

def tbl_row(slide, l, t, w, cells, idx):
    fc = CLG if idx % 2 == 0 else CW
    rect(slide, l, t, w, 0.45, fc)
    for val, cx, cw, al, sz, clr, bold in cells:
        txt(slide, l+cx, t+0.05, cw, 0.37, str(val),
            sz=sz, bold=bold, clr=clr, align=al)

# ══════════════════════════════════════════════════════
# 4. 슬라이드 생성
# ══════════════════════════════════════════════════════
prs = Presentation()
prs.slide_width  = Cm(33.87)
prs.slide_height = Cm(19.05)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

bg = slide.background.fill
bg.solid()
bg.fore_color.rgb = CW

# ══ [A] 타이틀 바 ═══════════════════════════════════
rect(slide, 0, 0, 33.87, 2.2, CN)
txt(slide, 0.6, 0.18, 27, 1.1,
    "SB30  부하 해리력 시험 부적합 발생 현황 분석",
    sz=22, bold=True, clr=CW)
txt(slide, 0.6, 1.38, 28, 0.65,
    "SB30 Buckle Release Force Test Analysis   ·   2017 – 2026 (10개년)",
    sz=9.5, clr=CLB)
txt(slide, 27.5, 0.55, 5.8, 0.8,
    "June  2026", sz=11, clr=CLB, align=PP_ALIGN.RIGHT)

# ══ [B] KPI 카드 4개 ════════════════════════════════
kpis = [
    ("총 시험 건수",  f"{total_n:,}건",    CN),
    ("전체 NG 건수",  f"{total_ng:,}건",   CR),
    ("전체 NG율",    f"{ng_rate:.1f}%",    CO),
    ("S1 상승 추세", "+0.86 N/년",         CG),
]
kw, kg, kx0 = 7.5, 0.49, 0.5
for i, (lbl, val, clr) in enumerate(kpis):
    kx = kx0 + i * (kw + kg)
    rect(slide, kx, 2.35, kw, 1.65, CLG, CMG)
    txt(slide, kx+0.15, 2.42, kw-0.3, 0.55, lbl,
        sz=8.5, clr=CDG, align=PP_ALIGN.CENTER)
    txt(slide, kx+0.15, 2.85, kw-0.3, 0.9, val,
        sz=19, bold=True, clr=clr, align=PP_ALIGN.CENTER)

# ══ [C] 왼쪽: 연도별 추세 차트 (w=11.2cm) ══════════
LX, LY, LW = 0.4, 4.2, 11.2
section_hdr(slide, LX, LY, LW, "① 연도별 NG율(%) 및 S1 평균 추세")
chart_buf = make_trend_chart()
slide.shapes.add_picture(chart_buf,
                         Cm(LX), Cm(LY+0.58), Cm(LW), Cm(7.6))

# ══ [D] 중앙: 고위험 차종 + 품번 (w=10.5cm) ════════
MX, MY, MW = 11.85, 4.2, 10.5

# ── 차종 TOP 7
section_hdr(slide, MX, MY, MW, "② 고위험 차종  NG율 상위 7개")
hy = MY + 0.58
tbl_hdr(slide, MX, hy, MW, [
    ("차종",   0.1, 3.2, PP_ALIGN.LEFT),
    ("총건수", 3.4, 3.0, PP_ALIGN.CENTER),
    ("NG율(%)", 6.6, 3.7, PP_ALIGN.RIGHT),
])
hy += 0.48
for i, (_, row) in enumerate(df_car.head(7).iterrows()):
    ng = row['NG율(%)']
    clr = CR if ng >= 30 else CO if ng >= 15 else CDG
    tbl_row(slide, MX, hy, MW, [
        (row['차종'],          0.1, 3.2, PP_ALIGN.LEFT,   8.0, CDG,  False),
        (f"{int(row['총건수'])}건", 3.4, 3.0, PP_ALIGN.CENTER, 7.5, CDG, False),
        (f"{ng:.1f}%",        6.6, 3.7, PP_ALIGN.RIGHT,  8.5, clr,  True),
    ], i)
    hy += 0.45

# ── 품번 TOP 5
hy += 0.18
section_hdr(slide, MX, hy, MW, "③ 고위험 품번  NG율 상위 5개", CMB)
hy += 0.58
tbl_hdr(slide, MX, hy, MW, [
    ("품번",    0.1, 7.2, PP_ALIGN.LEFT),
    ("NG율(%)", 7.5, 2.8, PP_ALIGN.RIGHT),
])
hy += 0.48
for i, (_, row) in enumerate(df_pn.head(5).iterrows()):
    ng = row['NG율(%)']
    clr = CR if ng >= 50 else CO
    tbl_row(slide, MX, hy, MW, [
        (row['품번'], 0.1, 7.2, PP_ALIGN.LEFT,  7.5, CDG, False),
        (f"{ng:.1f}%", 7.5, 2.8, PP_ALIGN.RIGHT, 8.5, clr, True),
    ], i)
    hy += 0.45

# ══ [E] 오른쪽: SPC + AI 예측 (w=11.0cm) ═══════════
RX, RY, RW = 22.6, 4.2, 11.0

# ── SPC 관리한계
section_hdr(slide, RX, RY, RW, "④ SPC 관리한계  (2024~2026,  n = 2,761건)", CMB)
sy = RY + 0.58
tbl_hdr(slide, RX, sy, RW, [
    ("항목",    0.1, 7.2, PP_ALIGN.LEFT),
    ("기준값",  7.4, 3.4, PP_ALIGN.RIGHT),
])
sy += 0.48
spc_rows = [
    ("UCL  (관리 상한, +3σ)", f"{spc_ucl:.2f} N",    CR),
    ("경보 상한 (+2σ)",        f"{spc_warn_u:.2f} N",  CO),
    ("CL   (중심선)",          f"{spc_cl:.2f} N",      CB),
    ("경보 하한 (-2σ)",        f"{spc_warn_l:.2f} N",  CO),
    ("LCL  (관리 하한, -3σ)", f"{spc_lcl:.2f} N",     CG),
]
for i, (lbl, val, clr) in enumerate(spc_rows):
    tbl_row(slide, RX, sy, RW, [
        (lbl, 0.1, 7.2, PP_ALIGN.LEFT,  8.0, CDG, False),
        (val, 7.4, 3.4, PP_ALIGN.RIGHT, 9.0, clr, True),
    ], i)
    sy += 0.46

# ── AI 예측
sy += 0.2
section_hdr(slide, RX, sy, RW, "⑤ AI 예측  향후 S1 해리력 수준", CMB)
sy += 0.58
pred_rows = [
    ("연간 상승 트렌드",         "+0.86 N/년",                              CB),
    ("2026 잔여 (6~12월)",      "49.3 ~ 51.5 N  (월별 예측)",               CDG),
    ("2027년 예측  (90% 구간)", f"{pred_2027[0]:.1f} ~ {pred_2027[1]:.1f} N", CB),
    ("2028년 예측  (90% 구간)", f"{pred_2028[0]:.1f} ~ {pred_2028[1]:.1f} N", CMB),
    ("계절성 피크  (6 ~ 8월)",  "+0.6 ~ +1.0 N  ▲",                        CO),
    ("계절성 저점  (1 ~ 3월)",  "-0.5 ~ -1.1 N  ▼",                        CG),
]
for i, (lbl, val, clr) in enumerate(pred_rows):
    tbl_row(slide, RX, sy, RW, [
        (lbl, 0.1, 6.6, PP_ALIGN.LEFT,  8.0, CDG, False),
        (val, 6.8, 4.0, PP_ALIGN.RIGHT, 8.0, clr, True),
    ], i)
    sy += 0.46

# ══ [F] 하단 단계별 NG 요약 바 ══════════════════════
fy = 12.45
rect(slide, 0.4, fy, 33.0, 0.58, CN)
txt(slide, 0.6, fy+0.11, 7.5, 0.4,
    "⑥ 시험 단계별 NG 현황", sz=8.5, bold=True, clr=CW)
sx = 8.0
for _, row in df_stage.iterrows():
    ng = row['NG율(%)']
    if ng == 0:
        clr = CLB
    elif ng >= 30:
        clr = CR
    elif ng >= 10:
        clr = CO
    else:
        clr = CG
    label = f"{row['단계']}  {ng:.1f}%"
    txt(slide, sx, fy+0.11, 4.9, 0.4, label,
        sz=8.5, bold=(ng >= 10), clr=clr)
    sx += 4.9
    if sx > 33: break

# ══ [G] 구분선 + 푸터 ═══════════════════════════════
rect(slide, 0, 18.42, 33.87, 0.02, CMG)
rect(slide, 0, 18.44, 33.87, 0.61, CLG)
txt(slide, 0.5, 18.51, 22, 0.45,
    "AI 모델: RandomForest  |  학습 데이터 6,995건 (2017~2026)  "
    "|  NG 분류 정확도 89.85%  |  회귀 R² = 0.32",
    sz=7.5, clr=CDG)
txt(slide, 23, 18.51, 10.5, 0.45,
    "SPC 기준기간: 2024~2026  ·  n = 2,761건",
    sz=7.5, clr=CDG, align=PP_ALIGN.RIGHT)

# ══ 저장 ════════════════════════════════════════════
prs.save(PPT_FILE)
print(f"PPT 저장 완료: {PPT_FILE}")
